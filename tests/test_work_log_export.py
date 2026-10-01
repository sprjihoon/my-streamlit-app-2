"""작업일지 엑셀: 화면 검색과 같은 필터의 전체 건을 담는다. 임시 DB만 사용."""
from __future__ import annotations

import asyncio
import io
from datetime import datetime

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from backend.app.api.work_log import export_work_logs, get_work_logs
from logic.db import get_connection


def _insert(**kwargs):
    payload = {
        "날짜": "2026-09-05",
        "업체명": "틸리언",
        "분류": "하차",
        "단가": 30000,
        "수량": 2,
        "비고1": "",
        "작성자": "장지훈",
        "출처": "bot",
    }
    payload.update(kwargs)
    payload["합계"] = int(payload["단가"]) * int(payload["수량"])
    with get_connection() as con:
        cur = con.execute(
            """INSERT INTO work_log
               (날짜, 업체명, 분류, 단가, 수량, 합계, 비고1, 작성자, 저장시간, 출처)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                payload["날짜"], payload["업체명"], payload["분류"], payload["단가"],
                payload["수량"], payload["합계"], payload["비고1"], payload["작성자"],
                datetime.now().isoformat(), payload["출처"],
            ),
        )
        con.commit()
        return int(cur.lastrowid)


def _body(resp) -> bytes:
    iterator = resp.body_iterator

    async def collect():
        chunks = []
        async for chunk in iterator:
            chunks.append(chunk)
        return b"".join(chunks)

    return asyncio.run(collect())


def _sheet_rows(data: bytes, sheet: str) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, row)) for row in rows[1:]]


def _export(**kwargs):
    params = dict(
        start_date="2026-09-01",
        end_date="2026-09-30",
        vendor="틸리언",
        work_type="하차",
        author="장지훈",
        source="bot",
        format="excel",
    )
    params.update(kwargs)
    return asyncio.run(export_work_logs(**params))


def test_export_keeps_filtered_rows_beyond_page(isolated_runtime):
    _insert(비고1="첫째")
    _insert(날짜="2026-09-06", 비고1="둘째")
    _insert(업체명="팔로우미코스메틱", 비고1="다른업체")
    _insert(분류="입고", 비고1="다른작업")
    _insert(작성자="다른사람", 비고1="다른작성자")
    _insert(출처="excel", 비고1="다른출처")
    _insert(날짜="2026-08-31", 비고1="지난달")

    listed = asyncio.run(get_work_logs(
        period_from="2026-09-01",
        period_to="2026-09-30",
        vendor="틸리언",
        work_type="하차",
        author="장지훈",
        source="bot",
        limit=1,
        offset=0,
    ))
    assert listed["total"] == 2
    assert len(listed["logs"]) == 1

    resp = _export()
    assert "spreadsheetml" in (resp.media_type or "")
    data = _body(resp)
    assert data[:2] == b"PK"
    rows = _sheet_rows(data, "작업일지")
    assert len(rows) == 2
    assert {row["업체명"] for row in rows} == {"틸리언"}
    assert {row["분류"] for row in rows} == {"하차"}
    assert {row["작성자"] for row in rows} == {"장지훈"}
    assert {row["출처"] for row in rows} == {"봇"}
    assert {row["비고"] for row in rows} == {"첫째", "둘째"}
    vendors = _sheet_rows(data, "업체별 요약")
    assert vendors[0]["업체명"] == "틸리언"
    assert vendors[0]["건수"] == 2
    assert vendors[0]["수량"] == 4
    assert vendors[0]["금액"] == 120000
    assert vendors[-1]["업체명"] == "합계"


def test_export_empty_filter_is_404(isolated_runtime):
    _insert(업체명="팔로우미코스메틱")
    with pytest.raises(HTTPException) as exc:
        _export(vendor="없는업체")
    assert exc.value.status_code == 404
    assert exc.value.detail == "해당 조건의 작업일지가 없습니다."


def test_export_excel_rejects_too_many_rows(isolated_runtime, monkeypatch):
    import logic.work_log_excel as excelmod

    monkeypatch.setattr(excelmod, "EXCEL_LOG_LIMIT", 1)
    _insert(비고1="첫째")
    _insert(날짜="2026-09-06", 비고1="둘째")
    with pytest.raises(HTTPException) as exc:
        _export()
    assert exc.value.status_code == 400
