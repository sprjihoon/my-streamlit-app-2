"""0원 인보이스 저장과 청구서 엑셀 병합 칸 테두리."""
import asyncio
from datetime import date

import pandas as pd
from openpyxl import Workbook

from backend.app.api.calculate import calculate_invoice
from backend.app.api.invoices import _create_invoice_sheet
from backend.app.models.schemas import InvoiceCalculateRequest
from logic.db import get_connection


def _closed(cell) -> bool:
    border = cell.border
    return all(getattr(border, side).style for side in ("left", "right", "top", "bottom"))


def _row_with(ws, label: str) -> int:
    for row in ws.iter_rows(min_row=1, max_row=40, max_col=1):
        if row[0].value == label:
            return row[0].row
    raise AssertionError(f"라벨 없음: {label}")


def test_merged_invoice_cells_have_closed_borders():
    wb = Workbook()
    ws = wb.active
    items = pd.DataFrame([
        {"항목": "기본 출고비", "수량": 15, "단가": 900, "금액": 13500, "비고": ""},
    ])
    _create_invoice_sheet(
        ws,
        {
            "invoice_id": 1442,
            "vendor_name": "휠",
            "period_from": "2026-09-01",
            "confirmed_by": "",
        },
        items,
        {
            "company_name": "틸리언",
            "business_number": "766-55-00323",
            "address": "대구시",
            "business_type": "서비스",
            "business_item": "포장 및 충전업",
            "bank_name": "카카오뱅크",
            "account_holder": "장지훈",
            "account_number": "3333",
            "representative": "장지훈",
        },
    )

    doc_row = _row_with(ws, "문서번호")
    for col in (2, 3, 5, 6):
        assert _closed(ws.cell(doc_row, col))

    recv_row = _row_with(ws, "수신")
    for col in range(2, 7):
        assert _closed(ws.cell(recv_row, col))

    supplier_row = _row_with(ws, "공급자")
    for row in range(supplier_row, supplier_row + 3):
        assert _closed(ws.cell(row, 1))
    for col in range(3, 7):
        assert _closed(ws.cell(supplier_row + 1, col))

    pay_row = _row_with(ws, "지급기한")
    for col in range(2, 7):
        assert _closed(ws.cell(pay_row, col))


def test_zero_amount_invoice_is_saved(isolated_runtime):
    with get_connection() as con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                username TEXT,
                password_hash TEXT,
                nickname TEXT,
                is_admin INTEGER
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token TEXT PRIMARY KEY,
                user_id INTEGER
            );
            CREATE TABLE IF NOT EXISTS invoices (
                invoice_id INTEGER PRIMARY KEY AUTOINCREMENT,
                vendor_id INTEGER,
                period_from TEXT,
                period_to TEXT,
                total_amount REAL,
                status TEXT,
                created_at TEXT
            );
            CREATE TABLE IF NOT EXISTS invoice_items (
                item_id INTEGER PRIMARY KEY AUTOINCREMENT,
                invoice_id INTEGER,
                item_name TEXT,
                qty REAL,
                unit_price REAL,
                amount REAL,
                remark TEXT
            );
            """
        )
        for col in ("barcode_f", "void_f", "pp_bag_f", "video_out_f", "video_ret_f"):
            con.execute(f"ALTER TABLE vendors ADD COLUMN {col} TEXT DEFAULT 'NO'")
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS kpost_ret (
                수취인명 TEXT,
                배달일자 TEXT,
                우편물부피 REAL
            )
            """
        )
        con.execute(
            "INSERT INTO vendors (vendor, name) VALUES ('제로상점', '제로상점')"
        )
        con.execute(
            "INSERT INTO users (username, password_hash, nickname, is_admin) VALUES ('a','x','테스터',1)"
        )
        con.execute("INSERT INTO sessions (token, user_id) VALUES ('tok-zero', 1)")
        con.commit()

    req = InvoiceCalculateRequest(
        vendor="제로상점",
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 30),
        include_basic_shipping=False,
        include_courier_fee=False,
        include_inbound_fee=False,
        include_remote_fee=False,
        include_worklog=False,
        include_combined_fee=False,
    )
    result = asyncio.run(calculate_invoice(req, token="tok-zero"))

    assert result.success is True
    assert result.total_amount == 0
    assert result.items == []
    assert result.invoice_id is not None

    with get_connection() as con:
        row = con.execute(
            "SELECT total_amount, status, vendor_id FROM invoices WHERE invoice_id=?",
            (result.invoice_id,),
        ).fetchone()
        item_count = con.execute(
            "SELECT COUNT(*) FROM invoice_items WHERE invoice_id=?",
            (result.invoice_id,),
        ).fetchone()[0]

    assert row[0] == 0
    assert row[1] == "미확정"
    assert item_count == 0
