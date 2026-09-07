# -*- coding: utf-8 -*-
"""수선작업일지 엑셀 보고. 업체·기간 필터 결과와 작업 사진을 포함한다."""

from __future__ import annotations

import io
from collections import defaultdict
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXCEL_LOG_LIMIT = 400
TEXT_HEADERS = ["날짜", "업체명", "제품명", "옵션", "바코드", "불량명", "작업", "수량", "비용", "비고", "작성자"]
PHOTO_PX = 84
ROW_HEIGHT = 66

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUMMARY_FILL = PatternFill("solid", fgColor="E8EEF6")
THIN = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def _existing_photos(log: dict) -> list[tuple[str, Path]]:
    out = []
    for label, path in log.get("photos") or []:
        if path and getattr(path, "is_file", None) and path.is_file():
            out.append((label, path))
    return out


def _fit_image(path: Path) -> Optional[XLImage]:
    try:
        img = XLImage(str(path))
    except Exception:
        return None
    width = img.width or PHOTO_PX
    height = img.height or PHOTO_PX
    ratio = min(PHOTO_PX / width, PHOTO_PX / height)
    img.width = max(1, int(width * ratio))
    img.height = max(1, int(height * ratio))
    return img


def _style_header(ws, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(columns)}1"


def _summary_rows(logs: list[dict]) -> list[tuple[Any, int, int, int]]:
    grouped: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0])
    for log in logs:
        name = (log.get("업체명") or "미지정").strip() or "미지정"
        bucket = grouped[name]
        bucket[0] += 1
        try:
            bucket[1] += int(log.get("수량") or 0)
        except (TypeError, ValueError):
            pass
        try:
            bucket[2] += int(log.get("비용") or 0)
        except (TypeError, ValueError):
            pass
    rows = [(name, *vals) for name, vals in grouped.items()]
    rows.sort(key=lambda r: (-r[3], r[0]))
    total = ("합계", sum(r[1] for r in rows), sum(r[2] for r in rows), sum(r[3] for r in rows))
    return [*rows, total]


def create_repair_log_xlsx(logs: list[dict]) -> bytes:
    """필터된 수선일지와 작업 사진이 들어간 xlsx 바이트를 만든다."""
    photo_lists = [_existing_photos(log) for log in logs]
    max_photos = max((len(p) for p in photo_lists), default=0)
    photo_headers = [f"사진{i}" for i in range(1, max_photos + 1)] if max_photos else ["사진"]
    headers = TEXT_HEADERS + photo_headers

    wb = Workbook()
    ws = wb.active
    ws.title = "수선일지"
    ws.append(headers)
    _style_header(ws, len(headers))

    for idx, log in enumerate(logs):
        row = idx + 2
        for col, key in enumerate(TEXT_HEADERS, start=1):
            cell = ws.cell(row, col, log.get(key) if log.get(key) is not None else "")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = THIN
        photos = photo_lists[idx]
        ws.row_dimensions[row].height = ROW_HEIGHT if photos else 18
        if not photos and max_photos == 0:
            ws.cell(row, len(TEXT_HEADERS) + 1, "")
        for offset, (_label, path) in enumerate(photos):
            col = len(TEXT_HEADERS) + 1 + offset
            ws.cell(row, col).border = THIN
            img = _fit_image(path)
            if img is None:
                ws.cell(row, col, "사진 오류")
                continue
            img.anchor = f"{get_column_letter(col)}{row}"
            ws.add_image(img)

    widths = [12, 16, 16, 12, 16, 12, 12, 8, 12, 18, 10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for i in range(len(TEXT_HEADERS) + 1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    summary = wb.create_sheet("업체별 요약")
    summary.append(["업체명", "건수", "수량", "금액"])
    _style_header(summary, 4)
    for name, count, qty, amount in _summary_rows(logs):
        summary.append([name, count, qty, amount])
    for col in range(1, 5):
        summary.column_dimensions[get_column_letter(col)].width = 16

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
