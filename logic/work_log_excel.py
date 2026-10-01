# -*- coding: utf-8 -*-
"""작업일지 엑셀 보고. 수선일지 보고와 같이 필터 결과와 업체별 요약을 담는다."""

from __future__ import annotations

import io
from collections import defaultdict
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXCEL_LOG_LIMIT = 5000
TEXT_HEADERS = ["날짜", "업체명", "분류", "수량", "단가", "합계", "비고", "작성자", "출처"]

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def _style_header(ws, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(columns)}1"


def _as_int(value) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _summary_rows(logs: list[dict]) -> list[tuple[Any, int, int, int]]:
    grouped: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0])
    for log in logs:
        name = (log.get("업체명") or "미지정").strip() or "미지정"
        bucket = grouped[name]
        bucket[0] += 1
        bucket[1] += _as_int(log.get("수량"))
        bucket[2] += _as_int(log.get("합계"))
    rows = [(name, *vals) for name, vals in grouped.items()]
    rows.sort(key=lambda r: (-r[3], r[0]))
    total = ("합계", sum(r[1] for r in rows), sum(r[2] for r in rows), sum(r[3] for r in rows))
    return [*rows, total]


def create_work_log_xlsx(logs: list[dict]) -> bytes:
    """필터된 작업일지가 들어간 xlsx 바이트를 만든다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "작업일지"
    ws.append(TEXT_HEADERS)
    _style_header(ws, len(TEXT_HEADERS))

    for idx, log in enumerate(logs):
        row = idx + 2
        for col, key in enumerate(TEXT_HEADERS, start=1):
            value = log.get(key)
            cell = ws.cell(row, col, value if value is not None else "")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = THIN
        ws.row_dimensions[row].height = 18

    widths = [12, 16, 14, 8, 12, 12, 18, 10, 10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    summary = wb.create_sheet("업체별 요약")
    summary.append(["업체명", "건수", "수량", "금액"])
    _style_header(summary, 4)
    for name, count, qty, amount in _summary_rows(logs):
        summary.append([name, count, qty, amount])
    for col in range(1, 5):
        summary.column_dimensions[get_column_letter(col)].width = 16

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
