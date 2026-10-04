"""Invoice worksheet layout shared by batch export."""
import pandas as pd
def _apply_range_border(ws, cell_range: str, border) -> None:
    """병합 범위의 모든 칸에 테두리를 그린다.

    openpyxl은 좌상단 칸만 칠하면 오른쪽·아래 선이 빠진다.
    """
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            ws.cell(row=row, column=col).border = border


def _create_invoice_sheet(ws, invoice_data: dict, items_df, company_info: dict):
    """인보이스 시트 생성 헬퍼 함수 (PDF와 동일한 양식)"""
    from datetime import datetime
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    
    # 스타일 정의
    title_font = Font(name='맑은 고딕', size=18, bold=True)
    header_font = Font(name='맑은 고딕', size=10, bold=True)
    body_font = Font(name='맑은 고딕', size=9)
    small_font = Font(name='맑은 고딕', size=8)
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    gray_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    light_gray_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    # 데이터 추출
    invoice_id = invoice_data['invoice_id']
    vendor_name = invoice_data['vendor_name']
    period_from = invoice_data['period_from']
    confirmed_by = invoice_data.get('confirmed_by', '')
    
    company_name = company_info.get('company_name', '')
    business_number = company_info.get('business_number', '')
    address = company_info.get('address', '')
    business_type = company_info.get('business_type', '')
    business_item = company_info.get('business_item', '')
    bank_name = company_info.get('bank_name', '')
    account_holder = company_info.get('account_holder', '')
    account_number = company_info.get('account_number', '')
    representative = company_info.get('representative', '')
    
    # 청구일자
    invoice_date = datetime.now().strftime("%Y-%m-%d")
    
    # 건명 생성
    period_str = period_from[:7].replace("-", "년 ") + "월" if period_from else ""
    title = f"{period_str} 풀필먼트 서비스 대금"
    
    # 수신자
    recipient_name = f"{vendor_name} 대표님 귀하"
    
    # 문서번호
    doc_number = f"{invoice_id:05d}-{invoice_date.replace('-', '')[:6]}"
    
    # 지급기한
    if period_from:
        try:
            from dateutil.relativedelta import relativedelta
            period_dt = datetime.strptime(period_from[:10], "%Y-%m-%d")
            next_month = period_dt + relativedelta(months=1)
            payment_deadline = f"{next_month.year}년 {next_month.month:02d}월 05일"
        except:
            payment_deadline = ""
    else:
        payment_deadline = ""
    
    current_row = 1
    
    # 1. 제목
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = "물류대행 서비스 대금청구서"
    ws[f'A{current_row}'].font = title_font
    ws[f'A{current_row}'].alignment = center_align
    current_row += 2
    
    # 2. 문서번호/청구일자
    ws[f'A{current_row}'] = "문서번호"
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].fill = gray_fill
    ws[f'A{current_row}'].border = thin_border
    ws.merge_cells(f'B{current_row}:C{current_row}')
    ws[f'B{current_row}'] = doc_number
    ws[f'B{current_row}'].font = body_font
    _apply_range_border(ws, f'B{current_row}:C{current_row}', thin_border)
    ws[f'D{current_row}'] = "청구일자"
    ws[f'D{current_row}'].font = header_font
    ws[f'D{current_row}'].fill = gray_fill
    ws[f'D{current_row}'].border = thin_border
    ws.merge_cells(f'E{current_row}:F{current_row}')
    ws[f'E{current_row}'] = invoice_date
    ws[f'E{current_row}'].font = body_font
    _apply_range_border(ws, f'E{current_row}:F{current_row}', thin_border)
    current_row += 1
    
    # 3. 수신/건명
    ws[f'A{current_row}'] = "수신"
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].fill = gray_fill
    ws[f'A{current_row}'].border = thin_border
    ws.merge_cells(f'B{current_row}:F{current_row}')
    ws[f'B{current_row}'] = recipient_name
    ws[f'B{current_row}'].font = body_font
    _apply_range_border(ws, f'B{current_row}:F{current_row}', thin_border)
    current_row += 1
    
    ws[f'A{current_row}'] = "건명"
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].fill = gray_fill
    ws[f'A{current_row}'].border = thin_border
    ws.merge_cells(f'B{current_row}:F{current_row}')
    ws[f'B{current_row}'] = title
    ws[f'B{current_row}'].font = body_font
    _apply_range_border(ws, f'B{current_row}:F{current_row}', thin_border)
    current_row += 1
    
    # 4. 공급자 정보
    ws[f'A{current_row}'] = "공급자"
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].fill = gray_fill
    ws[f'A{current_row}'].border = thin_border
    ws.merge_cells(f'A{current_row}:A{current_row+2}')
    _apply_range_border(ws, f'A{current_row}:A{current_row+2}', thin_border)
    
    ws[f'B{current_row}'] = "사업자번호"
    ws[f'B{current_row}'].font = small_font
    ws[f'B{current_row}'].fill = light_gray_fill
    ws[f'B{current_row}'].border = thin_border
    ws[f'C{current_row}'] = business_number
    ws[f'C{current_row}'].font = small_font
    ws[f'C{current_row}'].border = thin_border
    ws[f'D{current_row}'] = "상호"
    ws[f'D{current_row}'].font = small_font
    ws[f'D{current_row}'].fill = light_gray_fill
    ws[f'D{current_row}'].border = thin_border
    ws.merge_cells(f'E{current_row}:F{current_row}')
    ws[f'E{current_row}'] = company_name
    ws[f'E{current_row}'].font = small_font
    _apply_range_border(ws, f'E{current_row}:F{current_row}', thin_border)
    current_row += 1
    
    ws[f'B{current_row}'] = "소재지"
    ws[f'B{current_row}'].font = small_font
    ws[f'B{current_row}'].fill = light_gray_fill
    ws[f'B{current_row}'].border = thin_border
    ws.merge_cells(f'C{current_row}:F{current_row}')
    ws[f'C{current_row}'] = address
    ws[f'C{current_row}'].font = small_font
    _apply_range_border(ws, f'C{current_row}:F{current_row}', thin_border)
    current_row += 1
    
    ws[f'B{current_row}'] = "업태"
    ws[f'B{current_row}'].font = small_font
    ws[f'B{current_row}'].fill = light_gray_fill
    ws[f'B{current_row}'].border = thin_border
    ws[f'C{current_row}'] = business_type
    ws[f'C{current_row}'].font = small_font
    ws[f'C{current_row}'].border = thin_border
    ws[f'D{current_row}'] = "종목"
    ws[f'D{current_row}'].font = small_font
    ws[f'D{current_row}'].fill = light_gray_fill
    ws[f'D{current_row}'].border = thin_border
    ws.merge_cells(f'E{current_row}:F{current_row}')
    ws[f'E{current_row}'] = business_item
    ws[f'E{current_row}'].font = small_font
    _apply_range_border(ws, f'E{current_row}:F{current_row}', thin_border)
    current_row += 2
    
    # 5. 항목 테이블 헤더
    headers = ["No", "품명", "수량", "단가", "금액", "비고"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = gray_fill
        cell.border = thin_border
        cell.alignment = center_align
    current_row += 1
    
    # 항목 데이터
    subtotal = 0
    for idx, (_, row) in enumerate(items_df.iterrows(), 1):
        qty = int(row.get('수량', row.get('qty', 0))) if pd.notna(row.get('수량', row.get('qty'))) else 0
        unit_price = int(row.get('단가', row.get('unit_price', 0))) if pd.notna(row.get('단가', row.get('unit_price'))) else 0
        amount = int(row.get('금액', row.get('amount', qty * unit_price))) if pd.notna(row.get('금액', row.get('amount'))) else qty * unit_price
        item_name = str(row.get('항목', row.get('item_name', '')))
        remark = str(row.get('비고', row.get('remark', ''))) if pd.notna(row.get('비고', row.get('remark'))) else ""
        subtotal += amount
        
        ws.cell(row=current_row, column=1, value=idx).border = thin_border
        ws.cell(row=current_row, column=1).alignment = center_align
        ws.cell(row=current_row, column=2, value=item_name).border = thin_border
        ws.cell(row=current_row, column=3, value=f"{qty:,}" if qty else "").border = thin_border
        ws.cell(row=current_row, column=3).alignment = right_align
        ws.cell(row=current_row, column=4, value=f"{unit_price:,}" if unit_price else "").border = thin_border
        ws.cell(row=current_row, column=4).alignment = right_align
        ws.cell(row=current_row, column=5, value=f"{amount:,}" if amount else "").border = thin_border
        ws.cell(row=current_row, column=5).alignment = right_align
        ws.cell(row=current_row, column=6, value=remark).border = thin_border
        
        for col in range(1, 7):
            ws.cell(row=current_row, column=col).font = body_font
        current_row += 1
    
    current_row += 1
    
    # 6. 합계
    vat = int(subtotal * 0.1)
    total = subtotal + vat
    
    ws[f'A{current_row}'] = "합계 금액"
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].fill = gray_fill
    ws[f'A{current_row}'].border = thin_border
    ws[f'B{current_row}'] = f"₩ {subtotal:,}"
    ws[f'B{current_row}'].font = body_font
    ws[f'B{current_row}'].border = thin_border
    ws[f'B{current_row}'].alignment = right_align
    ws[f'C{current_row}'] = "부가세"
    ws[f'C{current_row}'].font = header_font
    ws[f'C{current_row}'].fill = gray_fill
    ws[f'C{current_row}'].border = thin_border
    ws[f'D{current_row}'] = f"₩ {vat:,}"
    ws[f'D{current_row}'].font = body_font
    ws[f'D{current_row}'].border = thin_border
    ws[f'D{current_row}'].alignment = right_align
    ws[f'E{current_row}'] = "청구금액"
    ws[f'E{current_row}'].font = header_font
    ws[f'E{current_row}'].fill = gray_fill
    ws[f'E{current_row}'].border = thin_border
    ws[f'F{current_row}'] = f"₩ {total:,}"
    ws[f'F{current_row}'].font = Font(name='맑은 고딕', size=11, bold=True)
    ws[f'F{current_row}'].border = thin_border
    ws[f'F{current_row}'].alignment = right_align
    current_row += 2
    
    # 7. 지급기한/계좌정보
    ws[f'A{current_row}'] = "지급기한"
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].fill = gray_fill
    ws[f'A{current_row}'].border = thin_border
    ws.merge_cells(f'B{current_row}:F{current_row}')
    ws[f'B{current_row}'] = payment_deadline
    ws[f'B{current_row}'].font = body_font
    _apply_range_border(ws, f'B{current_row}:F{current_row}', thin_border)
    current_row += 1
    
    ws[f'A{current_row}'] = "계좌정보"
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].fill = gray_fill
    ws[f'A{current_row}'].border = thin_border
    ws.merge_cells(f'B{current_row}:F{current_row}')
    ws[f'B{current_row}'] = f"{bank_name}  {account_number}  {account_holder}"
    ws[f'B{current_row}'].font = body_font
    _apply_range_border(ws, f'B{current_row}:F{current_row}', thin_border)
    current_row += 3
    
    # 8. 하단 - 위와 같이 청구합니다
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = "위와 같이 청구합니다."
    ws[f'A{current_row}'].font = header_font
    ws[f'A{current_row}'].alignment = center_align
    current_row += 2
    
    # 날짜 (한국어 형식)
    try:
        dt = datetime.strptime(invoice_date, "%Y-%m-%d")
        weekdays = ['월', '화', '수', '목', '금', '토', '일']
        date_str = f"{dt.year}년 {dt.month:02d}월 {dt.day:02d}일 {weekdays[dt.weekday()]}요일"
    except:
        date_str = invoice_date
    
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = date_str
    ws[f'A{current_row}'].font = body_font
    ws[f'A{current_row}'].alignment = center_align
    current_row += 2
    
    # 회사명
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = company_name
    ws[f'A{current_row}'].font = title_font
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # 담당자/대표자 정보
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = f"담당: {confirmed_by or '-'}  /  대표: {representative or '-'}"
    ws[f'A{current_row}'].font = small_font
    ws[f'A{current_row}'].alignment = center_align
    
    # 열 너비 조정
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 25
